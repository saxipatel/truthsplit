"""Read CSV / XLSX / XLS files into a header plus rows of raw cell values.

Nothing here interprets values; that is normalization's job. Readers only enforce structure: a known file
type whose signature matches its extension, one header row with unique names, and a row limit.
"""

import csv
import logging
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from app.models.sources import FileType

log = logging.getLogger(__name__)

Cell = str | int | float | Decimal | bool | datetime | date | time | None

ALLOWED_ENCODINGS = ("utf-8", "cp1252", "latin-1")
CSV_DELIMITERS = ",;\t|"

_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_MAX_XLSX_MEMBERS = 10_000
_MAX_XLSX_UNCOMPRESSED = 1024 * 1024 * 1024


class IngestionError(ValueError):
    """The file cannot be read as a table. The message is safe to show to the user."""


class UnsupportedFileError(IngestionError):
    pass


class EncodingError(IngestionError):
    """CSV bytes do not decode with the chosen encoding; another encoding may work."""


@dataclass
class RawTable:
    header: list[str]
    rows: list[list[Cell]] = field(default_factory=list)
    # Spreadsheet row number of each data row (the file's first row is 1). Blank rows are skipped, so
    # these are not always consecutive.
    row_numbers: list[int] = field(default_factory=list)
    # CSV rows must match the header width exactly; spreadsheet rows may omit trailing empty cells.
    strict_width: bool = False
    blank_rows_skipped: int = 0
    truncated: bool = False


def detect_file_type(filename: str, head: bytes) -> FileType:
    """Use the extension to choose a parser, and the leading bytes to check the file really is that type."""
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        if head.startswith((_ZIP_MAGIC, _OLE_MAGIC)) or b"\x00" in head:
            raise UnsupportedFileError("File has a .csv extension but contains binary data.")
        return FileType.CSV
    if ext == ".xlsx":
        if not head.startswith(_ZIP_MAGIC):
            raise UnsupportedFileError("File has a .xlsx extension but is not an Excel workbook.")
        return FileType.XLSX
    if ext == ".xls":
        if not head.startswith(_OLE_MAGIC):
            raise UnsupportedFileError("File has a .xls extension but is not a legacy Excel workbook.")
        return FileType.XLS
    raise UnsupportedFileError(f"Unsupported file type {ext or '(none)'}. Use .csv, .xlsx or .xls.")


def check_xlsx_container(path: Path) -> None:
    """Reject zip bombs, non-workbook archives and macro-enabled workbooks renamed to .xlsx."""
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise IngestionError("The .xlsx file is corrupt or not a zip archive.") from exc
    names = {i.filename for i in infos}
    if len(infos) > _MAX_XLSX_MEMBERS or sum(i.file_size for i in infos) > _MAX_XLSX_UNCOMPRESSED:
        raise IngestionError("The .xlsx file expands to more data than allowed.")
    if "xl/workbook.xml" not in names:
        raise IngestionError("The .xlsx archive does not contain a workbook.")
    if "xl/vbaProject.bin" in names:
        raise IngestionError("Macro-enabled workbooks are not accepted.")


def list_sheets(path: Path, file_type: FileType) -> list[str]:
    if file_type is FileType.CSV:
        return []
    if file_type is FileType.XLSX:
        wb = _open_xlsx(path)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    book = _open_xls(path)
    try:
        return book.sheet_names()
    finally:
        book.release_resources()


def read_table(
    path: Path,
    file_type: FileType,
    *,
    sheet: str | None = None,
    encoding: str = "utf-8",
    max_rows: int,
    preview: bool = False,
) -> RawTable:
    """Read the whole table. More than `max_rows` data rows is an error unless `preview` is set,
    in which case reading stops at `max_rows` and the table is marked truncated."""
    if file_type is FileType.CSV:
        if sheet:
            raise IngestionError("CSV files have no worksheets.")
        rows = _csv_rows(path, encoding)
        table = RawTable(header=[], strict_width=True)
    elif file_type is FileType.XLSX:
        rows = _xlsx_rows(path, sheet)
        table = RawTable(header=[])
    else:
        rows = _xls_rows(path, sheet)
        table = RawTable(header=[])

    header_row: list[Cell] | None = None
    try:
        for row_number, row in enumerate(rows, start=1):
            if header_row is None:
                if _is_blank(row):
                    continue  # leading blank rows before the header
                header_row = row
                table.header = _validate_header(row)
                continue
            if _is_blank(row):
                table.blank_rows_skipped += 1
                continue
            if len(table.rows) == max_rows:
                if preview:
                    table.truncated = True
                    break
                raise IngestionError(
                    f"The file has more than {max_rows:,} data rows, which is the import limit."
                )
            table.rows.append(row)
            table.row_numbers.append(row_number)
    finally:
        rows.close()  # releases the workbook even when reading stops early

    if header_row is None:
        raise IngestionError("The file is empty.")
    if not table.rows:
        raise IngestionError("The file has a header row but no data rows.")
    return table


def _validate_header(row: list[Cell]) -> list[str]:
    # Trailing empty header cells are common in spreadsheets; ignore them.
    cells = list(row)
    while cells and _is_empty(cells[-1]):
        cells.pop()
    header: list[str] = []
    for i, cell in enumerate(cells, start=1):
        name = "" if cell is None else str(cell).strip()
        if not name:
            raise IngestionError(f"Column {i} has an empty header.")
        if name in header:
            raise IngestionError(f"Duplicate column header {name!r}.")
        header.append(name)
    return header


def _is_empty(cell: Cell) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _is_blank(row: list[Cell]) -> bool:
    return all(_is_empty(c) for c in row)


def _csv_rows(path: Path, encoding: str):
    if encoding not in ALLOWED_ENCODINGS:
        raise IngestionError(
            f"Unsupported encoding {encoding!r}. Choose one of {', '.join(ALLOWED_ENCODINGS)}."
        )
    with path.open("rb") as fh:
        sample = fh.read(64 * 1024).decode(encoding, errors="ignore")
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    # utf-8-sig strips a byte-order mark if present.
    py_encoding = "utf-8-sig" if encoding == "utf-8" else encoding
    try:
        with path.open(newline="", encoding=py_encoding) as fh:
            reader = csv.reader(fh, delimiter=delimiter, strict=True)
            yield from reader
    except UnicodeDecodeError as exc:
        raise EncodingError(
            f"The file is not valid {encoding} text (byte offset {exc.start}). Try a different encoding."
        ) from exc
    except csv.Error as exc:
        raise IngestionError(f"Malformed CSV: {exc}.") from exc


def _open_xlsx(path: Path):
    import openpyxl  # defusedxml is installed, so openpyxl parses the workbook XML with it

    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many unrelated types for bad files
        log.info("xlsx open failed", extra={"error": f"{type(exc).__name__}: {exc}"})
        raise IngestionError("The .xlsx file could not be read as a workbook.") from exc


def _xlsx_rows(path: Path, sheet: str | None):
    wb = _open_xlsx(path)
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise IngestionError(f"Worksheet {sheet!r} not found.")
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def _open_xls(path: Path):
    try:
        import xlrd
    except ImportError as exc:
        raise UnsupportedFileError("XLS support requires the optional 'xlrd' package.") from exc
    try:
        return xlrd.open_workbook(path, on_demand=True)
    except xlrd.XLRDError as exc:
        raise IngestionError(f"The .xls file could not be read: {exc}.") from exc


def _xls_rows(path: Path, sheet: str | None):
    import xlrd

    book = _open_xls(path)
    try:
        if sheet is None:
            ws = book.sheet_by_index(0)
        elif sheet in book.sheet_names():
            ws = book.sheet_by_name(sheet)
        else:
            raise IngestionError(f"Worksheet {sheet!r} not found.")
        for r in range(ws.nrows):
            row: list[Cell] = []
            for c in ws.row(r):
                if c.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(c.value))
                elif c.ctype == xlrd.XL_CELL_ERROR:
                    row.append(xlrd.error_text_from_code.get(c.value, "#ERR"))
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(c.value)
            yield row
    finally:
        book.release_resources()
