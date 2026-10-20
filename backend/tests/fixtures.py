"""Builders for small CSV / XLSX / XLS files used across tests."""

from pathlib import Path

import openpyxl


def write_csv(path: Path, text: str, encoding: str = "utf-8") -> Path:
    path.write_bytes(text.encode(encoding))
    return path


def write_xlsx(path: Path, sheets: dict[str, list[list]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def write_xls(path: Path, rows: list[list], sheet: str = "Sheet1") -> Path:
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet(sheet)
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if hasattr(value, "year"):
                ws.write(r, c, value, date_style)
            elif value is not None:
                ws.write(r, c, value)
    wb.save(str(path))
    return path
